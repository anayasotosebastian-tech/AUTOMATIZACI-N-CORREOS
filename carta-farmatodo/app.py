import os
import re
import zipfile
import smtplib
import subprocess
import traceback
from pathlib import Path
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from flask import Flask, render_template, request, jsonify, send_file

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'farmatodo-lealtad-2026')

UPLOAD_FOLDER = 'uploads'
OUTPUT_FOLDER = 'output'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# In-memory session storage (para envío de correos posterior)
sessions = {}

MESES_ES = {
    1: 'enero', 2: 'febrero', 3: 'marzo', 4: 'abril',
    5: 'mayo', 6: 'junio', 7: 'julio', 8: 'agosto',
    9: 'septiembre', 10: 'octubre', 11: 'noviembre', 12: 'diciembre'
}
MESES_SHORT = {
    1: 'ENE', 2: 'FEB', 3: 'MAR', 4: 'ABR', 5: 'MAY', 6: 'JUN',
    7: 'JUL', 8: 'AGO', 9: 'SEP', 10: 'OCT', 11: 'NOV', 12: 'DIC'
}


def date_to_spanish(d):
    if d is None:
        return ''
    if isinstance(d, str):
        try:
            d = datetime.strptime(d, '%Y-%m-%d')
        except Exception:
            return d
    try:
        return f"{d.day} de {MESES_ES[d.month]}, {d.year}"
    except Exception:
        return str(d)


def date_to_short(d):
    if d is None:
        return ''
    if isinstance(d, str):
        try:
            d = datetime.strptime(d, '%Y-%m-%d')
        except Exception:
            return d
    try:
        return f"{d.day:02d}/{MESES_SHORT[d.month]}/{str(d.year)[2:]}"
    except Exception:
        return str(d)


def date_to_table(d):
    if d is None:
        return ''
    if isinstance(d, str):
        try:
            d = datetime.strptime(d, '%Y-%m-%d')
        except Exception:
            return d
    try:
        return d.strftime('%d/%m/%Y')
    except Exception:
        return str(d)


def read_excel(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            rows.append(dict(zip(headers, row)))
    return rows


def group_by_lab(rows):
    labs = {}
    for row in rows:
        lab_key = str(row.get('ID NADRO', '') or '').strip()
        if not lab_key or lab_key.lower() == 'none':
            continue
        if lab_key not in labs:
            labs[lab_key] = {
                'id_nadro': lab_key,
                'lab_nadro': str(row.get('LAB NADRO', '') or '').strip(),
                'laboratorio': str(row.get('LABORATORIO', '') or '').strip(),
                'nombre_kam': str(row.get('NOMBRE KAM', '') or '').strip(),
                'puesto': str(row.get('PUESTO', '') or '').strip(),
                'correo': str(row.get('CORREO', '') or '').strip(),
                'copia_correo': str(row.get('COPIA CORREO', '') or '').strip(),
                'fecha_inicio': row.get('FECHA INICIO'),
                'fecha_fin': row.get('FECHA FIN'),
                'productos': []
            }
        labs[lab_key]['productos'].append({
            'sap': str(row.get('SAP', '') or '').strip(),
            'sku': str(row.get('SKU', '') or '').strip(),
            'producto': str(row.get('Producto', '') or '').strip(),
            'mecanica': str(row.get('Mecánica', '') or '').strip(),
            'limites': str(row.get('Límites', '') or '').strip(),
            'fecha_inicio': row.get('FECHA INICIO'),
            'fecha_fin': row.get('FECHA FIN'),
        })
    return labs


# ---------- Word document helpers ----------

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_run(para, text, size=11, bold=False, italic=False, color=None):
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def generate_word_doc(lab_data, doc_date, forma_recuperacion, output_path):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    fi_lab = lab_data['fecha_inicio']
    ff_lab = lab_data['fecha_fin']
    inicio_str = date_to_short(fi_lab)
    fin_str = date_to_short(ff_lab)

    # --- Fecha ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(p, doc_date)

    doc.add_paragraph()

    # --- Párrafo principal ---
    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p_main, 'Por medio del presente se confirma la participación del Proveedor ')
    add_run(p_main, f"{lab_data['id_nadro']} {lab_data['lab_nadro']}", bold=True)
    add_run(p_main, ' en el Plan de Lealtad de Grupo Farmatodo iniciando operaciones en mostradores afiliados a este programa a partir del ')
    add_run(p_main, inicio_str, bold=True)
    add_run(p_main, ' y cuya vigencia será extensiva hasta el ')
    add_run(p_main, fin_str, bold=True)
    add_run(p_main, ', con una prórroga mínima de 90 días para el cierre de ciclos.')

    doc.add_paragraph()

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p2, 'El tipo de promoción y/o beneficio y los productos participantes se encuentran definidos en el Anexo I.')

    doc.add_paragraph()

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p3, 'La mecánica de recuperación será:', bold=True)

    doc.add_paragraph()

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p4, 'Durante los primeros 10 (diez) días hábiles, el equipo de Plan de Lealtad '
               '(sebastian_anaya@nadro.com.mx) enviará el reporte detallado al laboratorio, '
               'este reporte será validado previamente por el Laboratorio, quien contará con '
               '10 (diez) días hábiles para dar el visto bueno para el pago de la recuperación '
               'a precio farmacia.')

    doc.add_paragraph()

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p5, 'Cualquier modificación al producto como baja, cambio de sku o exclusión para '
               'farmacias es necesario informar al equipo de Plan de Lealtad por correo '
               'electrónico (omar.enciso@rfp.mx, sebastian_anaya@nadro.com.mx, '
               'eltarjeton@rfp.mx, josep.jaramillo@rfp.mx), con una anticipación mínima de '
               '30 días antes a la fecha de aplicación, esto tomando en cuenta que las fechas '
               'para movimiento al sistema son los días 01 de cada mes.')

    doc.add_paragraph()

    p6 = doc.add_paragraph()
    add_run(p6, 'Anexo I.', bold=True)

    doc.add_paragraph()

    # --- Tabla OC / Costo 0 ---
    oc_table = doc.add_table(rows=2, cols=3)
    oc_table.style = 'Table Grid'
    oc_table.alignment = WD_TABLE_ALIGNMENT.LEFT

    header_cells = oc_table.rows[0].cells
    for i, txt in enumerate(['Forma de Recuperación', 'OC', 'Costo 0']):
        header_cells[i].text = txt
        header_cells[i].paragraphs[0].runs[0].bold = True
        header_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_bg(header_cells[i], '1F497D')
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    val_cells = oc_table.rows[1].cells
    val_cells[0].text = ''
    oc_val = 'X' if forma_recuperacion == 'OC' else ''
    costo_val = 'X' if forma_recuperacion == 'Costo 0' else ''
    for i, txt in enumerate(['', oc_val, costo_val]):
        val_cells[i].text = txt
        val_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if txt == 'X':
            val_cells[i].paragraphs[0].runs[0].bold = True
            val_cells[i].paragraphs[0].runs[0].font.size = Pt(14)

    oc_table.columns[0].width = Cm(6)
    oc_table.columns[1].width = Cm(3)
    oc_table.columns[2].width = Cm(3)

    doc.add_paragraph()
    doc.add_paragraph()

    # --- Firmas ---
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.style = 'Table Grid'

    for c in sig_table.rows[0].cells:
        c.text = '_' * 32
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    sig_table.rows[1].cells[0].text = 'Josep Jaramillo'
    sig_table.rows[1].cells[0].paragraphs[0].runs[0].bold = True
    sig_table.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    sig_table.rows[1].cells[1].text = lab_data['nombre_kam']
    sig_table.rows[1].cells[1].paragraphs[0].runs[0].bold = True
    sig_table.rows[1].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    sig_table.rows[2].cells[0].text = 'Gerente de Planes de Lealtad\nGrupo Farmatodo'
    sig_table.rows[2].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    sig_table.rows[2].cells[1].text = lab_data['puesto']
    sig_table.rows[2].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph()

    # --- Productos (Anexo I) ---
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_titulo, 'ANEXO I – PRODUCTOS PARTICIPANTES', bold=True, size=12)

    doc.add_paragraph()

    col_names = ['SAP', 'SKU', 'Producto', 'Mecánica', 'Límites', 'Vigencia']
    prod_table = doc.add_table(rows=1, cols=6)
    prod_table.style = 'Table Grid'

    hdr = prod_table.rows[0]
    for i, h in enumerate(col_names):
        hdr.cells[i].text = h
        hdr.cells[i].paragraphs[0].runs[0].bold = True
        hdr.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_bg(hdr.cells[i], '1F497D')
        hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for prod in lab_data['productos']:
        fi_p = prod.get('fecha_inicio') or fi_lab
        ff_p = prod.get('fecha_fin') or ff_lab
        vigencia = ''
        if fi_p and ff_p:
            vigencia = f"{date_to_table(fi_p)} – {date_to_table(ff_p)}"

        row = prod_table.add_row()
        values = [prod['sap'], prod['sku'], prod['producto'],
                  prod['mecanica'], prod['limites'], vigencia]
        for i, val in enumerate(values):
            row.cells[i].text = str(val)
            if row.cells[i].paragraphs[0].runs:
                row.cells[i].paragraphs[0].runs[0].font.size = Pt(9)

    col_widths = [Cm(1.4), Cm(3.2), Cm(5.5), Cm(1.8), Cm(1.8), Cm(3.3)]
    for i, w in enumerate(col_widths):
        for cell in prod_table.columns[i].cells:
            cell.width = w

    doc.save(output_path)


def convert_to_pdf(docx_path, output_dir):
    try:
        result = subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', output_dir, docx_path],
            capture_output=True, text=True, timeout=120
        )
        pdf_name = Path(docx_path).stem + '.pdf'
        pdf_path = os.path.join(output_dir, pdf_name)
        return pdf_path if os.path.exists(pdf_path) else None
    except FileNotFoundError:
        try:
            result = subprocess.run(
                ['soffice', '--headless', '--convert-to', 'pdf', '--outdir', output_dir, docx_path],
                capture_output=True, text=True, timeout=120
            )
            pdf_name = Path(docx_path).stem + '.pdf'
            pdf_path = os.path.join(output_dir, pdf_name)
            return pdf_path if os.path.exists(pdf_path) else None
        except Exception:
            return None
    except Exception as e:
        print(f"PDF error: {e}")
        return None


def safe_filename(name):
    return re.sub(r'[^\w\-]', '_', name)


def build_email_body(lab_data):
    primer_nombre = lab_data['nombre_kam'].split()[0] if lab_data['nombre_kam'] else 'estimado(a)'
    return (
        f"Buen día, {primer_nombre}, ¿cómo estás?\n\n"
        "Nos gustaría iniciar el proceso de renovación del Plan de Lealtad que mantenemos "
        "con ustedes. Para nosotros es muy importante continuar trabajando de manera conjunta, "
        "por lo que agradeceríamos su apoyo para actualizar la carta correspondiente y revisar "
        "si será posible mantener las condiciones actuales o si consideran necesario proponer "
        "algún ajuste.\n\n"
        "Quedamos atentos y en espera de tu apoyo con la revisión y firma de la carta de renovación.\n\n"
        "Agradecemos mucho su apoyo y el excelente trabajo que hemos realizado durante este periodo."
    )


def send_lab_email(smtp_cfg, lab_data, doc_date, word_path, pdf_path):
    nombre_lab = lab_data['lab_nadro']
    subject = f"SOLICITUD DE RENOVACIÓN PLAN DE LEALTAD - {nombre_lab} {doc_date}"

    msg = MIMEMultipart()
    msg['From'] = smtp_cfg['from_email']
    msg['To'] = lab_data['correo']
    msg['Subject'] = subject

    cc_list = [e.strip() for e in lab_data['copia_correo'].split(',') if e.strip()] if lab_data['copia_correo'] else []
    if cc_list:
        msg['Cc'] = ', '.join(cc_list)

    msg.attach(MIMEText(build_email_body(lab_data), 'plain', 'utf-8'))

    for fpath in [word_path, pdf_path]:
        if fpath and os.path.exists(fpath):
            with open(fpath, 'rb') as f:
                part = MIMEApplication(f.read(), Name=os.path.basename(fpath))
                part['Content-Disposition'] = f'attachment; filename="{os.path.basename(fpath)}"'
                msg.attach(part)

    port = int(smtp_cfg.get('port', 587))
    with smtplib.SMTP(smtp_cfg['host'], port) as server:
        server.ehlo()
        server.starttls()
        server.login(smtp_cfg['username'], smtp_cfg['password'])
        recipients = [lab_data['correo']] + cc_list
        server.sendmail(smtp_cfg['from_email'], recipients, msg.as_string())


# ---------- Routes ----------

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/process', methods=['POST'])
def process():
    try:
        excel_file = request.files.get('excel_file')
        doc_name = request.form.get('doc_name', 'CARTA_PLAN_LEALTAD').strip() or 'CARTA_PLAN_LEALTAD'
        doc_date_str = request.form.get('doc_date', datetime.now().strftime('%Y-%m-%d'))
        forma_rec = request.form.get('forma_recuperacion', 'OC')

        if not excel_file:
            return jsonify({'error': 'No se recibió ningún archivo Excel'}), 400

        excel_path = os.path.join(UPLOAD_FOLDER, 'upload_temp.xlsx')
        excel_file.save(excel_path)

        rows = read_excel(excel_path)
        if not rows:
            return jsonify({'error': 'El archivo Excel está vacío o no tiene datos válidos'}), 400

        labs = group_by_lab(rows)
        if not labs:
            return jsonify({'error': 'No se encontraron laboratorios válidos en el Excel'}), 400

        try:
            date_obj = datetime.strptime(doc_date_str, '%Y-%m-%d')
            doc_date_formatted = date_to_spanish(date_obj)
        except Exception:
            doc_date_formatted = doc_date_str

        session_id = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_dir = os.path.join(OUTPUT_FOLDER, session_id)
        os.makedirs(output_dir, exist_ok=True)

        results = []
        for lab_id, lab_data in labs.items():
            safe = safe_filename(lab_data['lab_nadro'])
            word_filename = f"{safe_filename(doc_name)}_{safe}.docx"
            word_path = os.path.join(output_dir, word_filename)

            generate_word_doc(lab_data, doc_date_formatted, forma_rec, word_path)

            pdf_path = convert_to_pdf(word_path, output_dir)
            pdf_filename = Path(word_filename).stem + '.pdf' if pdf_path else None

            results.append({
                'lab_id': lab_id,
                'lab_name': lab_data['lab_nadro'],
                'laboratorio': lab_data['laboratorio'],
                'kam': lab_data['nombre_kam'],
                'email': lab_data['correo'],
                'productos_count': len(lab_data['productos']),
                'word_file': word_filename,
                'pdf_file': pdf_filename,
                'session_id': session_id,
            })

        sessions[session_id] = {
            'labs': labs,
            'doc_date': doc_date_formatted,
            'doc_name': doc_name,
            'forma_rec': forma_rec,
            'output_dir': output_dir,
        }

        return jsonify({'success': True, 'results': results, 'session_id': session_id,
                        'total_labs': len(labs), 'total_productos': len(rows)})

    except Exception as e:
        return jsonify({'error': str(e), 'detail': traceback.format_exc()}), 500


@app.route('/download/<session_id>/<filename>')
def download_file(session_id, filename):
    filename = Path(filename).name  # security: no path traversal
    output_dir = os.path.join(OUTPUT_FOLDER, session_id)
    filepath = os.path.join(output_dir, filename)
    if not os.path.exists(filepath):
        return 'Archivo no encontrado', 404
    return send_file(os.path.abspath(filepath), as_attachment=True)


@app.route('/preview/<session_id>/<filename>')
def preview_file(session_id, filename):
    filename = Path(filename).name
    output_dir = os.path.join(OUTPUT_FOLDER, session_id)
    filepath = os.path.join(output_dir, filename)
    if not os.path.exists(filepath):
        return 'Archivo no encontrado', 404
    mimetype = 'application/pdf' if filename.endswith('.pdf') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    return send_file(os.path.abspath(filepath), mimetype=mimetype)


@app.route('/download_all/<session_id>')
def download_all(session_id):
    output_dir = os.path.join(OUTPUT_FOLDER, session_id)
    if not os.path.exists(output_dir):
        return 'Sesión no encontrada', 404
    zip_path = os.path.join(OUTPUT_FOLDER, f'cartas_{session_id}.zip')
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for f in os.listdir(output_dir):
            zf.write(os.path.join(output_dir, f), f)
    return send_file(os.path.abspath(zip_path), as_attachment=True,
                     download_name=f'Cartas_Plan_Lealtad_{session_id}.zip')


@app.route('/send_email', methods=['POST'])
def send_email_route():
    data = request.json or {}
    session_id = data.get('session_id')
    lab_id = data.get('lab_id')
    smtp_cfg = data.get('smtp_config', {})

    if session_id not in sessions:
        return jsonify({'error': 'Sesión no encontrada'}), 404

    sd = sessions[session_id]
    lab_data = sd['labs'].get(lab_id)
    if not lab_data:
        return jsonify({'error': 'Laboratorio no encontrado'}), 404

    safe = safe_filename(lab_data['lab_nadro'])
    word_filename = f"{safe_filename(sd['doc_name'])}_{safe}.docx"
    word_path = os.path.join(sd['output_dir'], word_filename)
    pdf_path = os.path.join(sd['output_dir'], Path(word_filename).stem + '.pdf')

    try:
        send_lab_email(smtp_cfg, lab_data, sd['doc_date'],
                       word_path if os.path.exists(word_path) else None,
                       pdf_path if os.path.exists(pdf_path) else None)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/send_all_emails', methods=['POST'])
def send_all_emails():
    data = request.json or {}
    session_id = data.get('session_id')
    smtp_cfg = data.get('smtp_config', {})

    if session_id not in sessions:
        return jsonify({'error': 'Sesión no encontrada'}), 404

    sd = sessions[session_id]
    sent = []
    errors = []

    for lab_id, lab_data in sd['labs'].items():
        safe = safe_filename(lab_data['lab_nadro'])
        word_filename = f"{safe_filename(sd['doc_name'])}_{safe}.docx"
        word_path = os.path.join(sd['output_dir'], word_filename)
        pdf_path = os.path.join(sd['output_dir'], Path(word_filename).stem + '.pdf')
        try:
            send_lab_email(smtp_cfg, lab_data, sd['doc_date'],
                           word_path if os.path.exists(word_path) else None,
                           pdf_path if os.path.exists(pdf_path) else None)
            sent.append(lab_data['lab_nadro'])
        except Exception as e:
            errors.append({'lab': lab_data['lab_nadro'], 'error': str(e)})

    return jsonify({'success': len(errors) == 0, 'sent': sent, 'errors': errors})


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
