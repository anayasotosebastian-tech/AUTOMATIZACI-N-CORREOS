import os
import io
import json
import copy
import subprocess
import tempfile
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, session
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX

app = Flask(__name__)
app.secret_key = "farmatodo-lealtad-2026"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(BASE_DIR, "plantilla.docx")
MATERIALES = os.path.join(BASE_DIR, "materiales.xlsx")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOADS_DIR, exist_ok=True)

# ── Load catalog once at startup ──────────────────────────────────────────────
def load_catalog():
    wb = openpyxl.load_workbook(MATERIALES, read_only=True, data_only=True)
    ws = wb["Hoja1"]
    rows = list(ws.iter_rows(values_only=True))
    # rows[0] = ('SAP','SKU','DESCRIPCION','ID LAB','LAB')
    sku_map = {}    # sku_str -> {sap, descripcion, id_lab, lab}
    lab_map = {}    # id_lab_str -> lab_name
    for row in rows[1:]:
        sap, sku, desc, id_lab, lab = row
        if sku is None:
            continue
        sku_str = str(int(sku)) if isinstance(sku, float) else str(sku)
        sap_str = str(int(sap)) if sap and isinstance(sap, (int, float)) else (str(sap) if sap else "-")
        entry = {"sap": sap_str, "descripcion": str(desc) if desc else "-", "id_lab": str(int(id_lab)) if id_lab else "-", "lab": str(lab) if lab else "-"}
        sku_map[sku_str] = entry
        if id_lab:
            id_lab_str = str(int(id_lab)) if isinstance(id_lab, float) else str(id_lab)
            if id_lab_str not in lab_map:
                lab_map[id_lab_str] = str(lab) if lab else "-"
    wb.close()
    return sku_map, lab_map

SKU_MAP, LAB_MAP = load_catalog()


# ── Helpers ───────────────────────────────────────────────────────────────────
def replace_yellow_runs(paragraph, new_text):
    """Replace all yellow-highlighted runs in a paragraph with new_text in the first yellow run, clearing the rest."""
    yellow_runs = [r for r in paragraph.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    if not yellow_runs:
        return
    # Put full replacement in first yellow run, clear others
    yellow_runs[0].text = new_text
    yellow_runs[0].font.highlight_color = None
    for r in yellow_runs[1:]:
        r.text = ""
        r.font.highlight_color = None


def build_document(data):
    """Build a filled Word document from template + form data. Returns Document object."""
    doc = Document(PLANTILLA)

    fecha_doc = data.get("fecha_doc", "")
    id_lab = data.get("id_lab", "")
    nombre_lab = data.get("nombre_lab", "")
    inicio_vigencia = data.get("inicio_vigencia", "")
    fin_vigencia = data.get("fin_vigencia", "")
    forma_recuperacion = data.get("forma_recuperacion", "")
    nombre_kam = data.get("nombre_kam", "")
    puesto = data.get("puesto", "")
    productos = data.get("productos", [])  # list of {sku, sap, producto, mecanica, limites}

    # ── Paragraph 2: fecha ────────────────────────────────────────────────────
    replace_yellow_runs(doc.paragraphs[2], fecha_doc)

    # ── Paragraph 4: id_lab + nombre_lab + inicio + fin ───────────────────────
    p4 = doc.paragraphs[4]
    yellow4 = [r for r in p4.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    # yellow4[0] = id_lab, yellow4[1] = nombre_lab, yellow4[2..4] = inicio, yellow4[5..7] = fin
    repl_values = [id_lab, nombre_lab, inicio_vigencia, "", "", fin_vigencia, "", ""]
    replacements = [id_lab, nombre_lab]
    # Parse inicio dd/MMM/yy
    def format_vigencia(date_str):
        """Convert YYYY-MM-DD to DD/MMM/YY in Spanish"""
        months = ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"]
        try:
            dt = datetime.strptime(date_str, "%Y-%m-%d")
            return f"{dt.day:02d}/{months[dt.month-1]}/{str(dt.year)[2:]}"
        except Exception:
            return date_str

    inicio_fmt = format_vigencia(inicio_vigencia)
    fin_fmt = format_vigencia(fin_vigencia)
    # inicio parts
    inicio_parts = inicio_fmt.split("/") if "/" in inicio_fmt else [inicio_fmt, "", ""]
    fin_parts = fin_fmt.split("/") if "/" in fin_fmt else [fin_fmt, "", ""]
    vigencia_replacements = inicio_parts + fin_parts  # 6 items

    all_replacements = [id_lab, nombre_lab] + vigencia_replacements
    for idx, run in enumerate(yellow4):
        val = all_replacements[idx] if idx < len(all_replacements) else ""
        run.text = val
        run.font.highlight_color = None

    # ── Table 0: forma de recuperación ────────────────────────────────────────
    t0 = doc.tables[0]
    cell00 = t0.rows[0].cells[0]
    for p in cell00.paragraphs:
        replace_yellow_runs(p, forma_recuperacion)

    # ── Table 1: KAM name, puesto, lab name ───────────────────────────────────
    t1 = doc.tables[1]
    # Row 1 col 1: nombre KAM
    for p in t1.rows[1].cells[1].paragraphs:
        replace_yellow_runs(p, nombre_kam)
    # Row 2 col 1: puesto
    for p in t1.rows[2].cells[1].paragraphs:
        replace_yellow_runs(p, puesto)
    # Row 3 col 1: nombre lab
    for p in t1.rows[3].cells[1].paragraphs:
        replace_yellow_runs(p, nombre_lab)

    # ── Table 2: productos ────────────────────────────────────────────────────
    t2 = doc.tables[2]
    # Keep header row, remove existing data rows
    header_row = t2.rows[0]
    # Remove all rows after header
    for row in t2.rows[1:]:
        tr = row._tr
        tr.getparent().remove(tr)

    # Re-add rows for each product
    for prod in productos:
        row = t2.add_row()
        vigencia_cell = f"{inicio_fmt} – {fin_fmt}"
        values = [
            prod.get("sap", "-"),
            prod.get("sku", "-"),
            prod.get("producto", "-"),
            prod.get("mecanica", "-"),
            prod.get("limites", "-"),
            vigencia_cell,
        ]
        for i, cell in enumerate(row.cells):
            cell.text = values[i] if i < len(values) else ""
            # Copy formatting from original header to match style
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = header_row.cells[0].paragraphs[0].runs[0].font.size if header_row.cells[0].paragraphs[0].runs else None

    return doc


def doc_to_pdf(doc, out_path):
    """Save doc to docx then convert to PDF using LibreOffice."""
    tmp_docx = out_path.replace(".pdf", ".docx")
    doc.save(tmp_docx)
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir",
         os.path.dirname(out_path), tmp_docx],
        capture_output=True, text=True, timeout=60
    )
    # libreoffice names the pdf same as docx
    expected_pdf = tmp_docx.replace(".docx", ".pdf")
    if os.path.exists(expected_pdf) and expected_pdf != out_path:
        os.rename(expected_pdf, out_path)
    return result.returncode == 0


# ── Routes ────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/lab")
def api_lab():
    id_lab = request.args.get("id", "").strip()
    name = LAB_MAP.get(id_lab, "")
    return jsonify({"nombre": name})


@app.route("/api/sku", methods=["POST"])
def api_sku():
    skus_raw = request.json.get("skus", [])
    results = []
    for sku in skus_raw:
        sku_str = str(sku).strip()
        entry = SKU_MAP.get(sku_str, None)
        if entry:
            results.append({
                "sku": sku_str,
                "sap": entry["sap"],
                "producto": entry["descripcion"],
                "found": True,
            })
        else:
            results.append({"sku": sku_str, "sap": "-", "producto": "-", "found": False})
    return jsonify(results)


@app.route("/api/preview", methods=["POST"])
def api_preview():
    data = request.json
    doc = build_document(data)
    # Save to temp file and convert to PDF
    uid = datetime.now().strftime("%Y%m%d%H%M%S%f")
    pdf_path = os.path.join(UPLOADS_DIR, f"preview_{uid}.pdf")
    docx_path = os.path.join(UPLOADS_DIR, f"preview_{uid}.docx")
    doc.save(docx_path)
    ok = doc_to_pdf(doc, pdf_path)
    if ok and os.path.exists(pdf_path):
        # Store paths in session for later use
        session["last_docx"] = docx_path
        session["last_pdf"] = pdf_path
        session["last_data"] = json.dumps(data)
        return jsonify({"ok": True, "pdf_url": f"/api/pdf/{uid}"})
    else:
        # fallback: return docx
        session["last_docx"] = docx_path
        session["last_data"] = json.dumps(data)
        return jsonify({"ok": False, "docx_url": f"/api/docx/{uid}"})


@app.route("/api/pdf/<uid>")
def serve_pdf(uid):
    path = os.path.join(UPLOADS_DIR, f"preview_{uid}.pdf")
    if os.path.exists(path):
        return send_file(path, mimetype="application/pdf")
    return "Not found", 404


@app.route("/api/docx/<uid>")
def serve_docx(uid):
    path = os.path.join(UPLOADS_DIR, f"preview_{uid}.docx")
    if os.path.exists(path):
        return send_file(path, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    return "Not found", 404


@app.route("/api/generate", methods=["POST"])
def api_generate():
    """Generate final docx + pdf and return download links."""
    data = request.json
    doc = build_document(data)
    uid = datetime.now().strftime("%Y%m%d%H%M%S%f")
    nombre_doc = data.get("nombre_documento", "convenio").replace(" ", "_")
    docx_path = os.path.join(UPLOADS_DIR, f"{nombre_doc}_{uid}.docx")
    pdf_path = os.path.join(UPLOADS_DIR, f"{nombre_doc}_{uid}.pdf")
    doc.save(docx_path)
    ok = doc_to_pdf(doc, pdf_path)
    resp = {
        "docx_url": f"/download/{os.path.basename(docx_path)}",
    }
    if ok and os.path.exists(pdf_path):
        resp["pdf_url"] = f"/download/{os.path.basename(pdf_path)}"
    return jsonify(resp)


@app.route("/download/<filename>")
def download_file(filename):
    path = os.path.join(UPLOADS_DIR, filename)
    if os.path.exists(path):
        return send_file(path, as_attachment=True)
    return "Not found", 404


@app.route("/api/email_body", methods=["POST"])
def api_email_body():
    data = request.json
    accion = data.get("accion", "RENOVACION")
    nombre_lab = data.get("nombre_lab", "")
    nombre_kam = data.get("nombre_kam", "")
    fecha_doc = data.get("fecha_doc", "")
    primer_nombre = nombre_kam.split()[0] if nombre_kam else "equipo"

    if accion == "RENOVACION":
        subject = f"SOLICITUD DE RENOVACIÓN PLAN DE LEALTAD - {nombre_lab} {fecha_doc}"
        body = f"""Buen día, {primer_nombre}, ¿cómo estás?

Nos gustaría iniciar el proceso de renovación del Plan de Lealtad que mantenemos con ustedes. Para nosotros es muy importante continuar trabajando de manera conjunta, por lo que agradeceríamos su apoyo para actualizar la carta correspondiente y revisar si será posible mantener las condiciones actuales o si consideran necesario proponer algún ajuste.

Quedamos atentos y en espera de tu apoyo con la revisión y firma de la carta de renovación.

Agradecemos mucho su apoyo y el excelente trabajo que hemos realizado durante este periodo."""
    else:
        subject = f"CARTA PLAN DE LEALTAD {nombre_lab} - GRUPO FT 2026"
        body = f"""Buen día, {primer_nombre}, esperando se encuentren con bien,

Les comparto la carta convenio con los productos revisados para el plan exclusivo de Farmatodo.

¿Me puedes apoyar con la firma del mismo? Por favor

Les agradecemos el apoyo y estamos seguros que tendremos excelentes resultados con su incorporación al Programa.

Cualquier duda o comentario quedamos al pendiente."""

    return jsonify({"subject": subject, "body": body})


if __name__ == "__main__":
    app.run(debug=True, port=5000)
