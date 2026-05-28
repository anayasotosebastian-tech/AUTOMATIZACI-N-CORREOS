import os
import io
import subprocess
import tempfile
import base64
from datetime import date, datetime
import openpyxl
import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(BASE_DIR, "plantilla.docx")
MATERIALES = os.path.join(BASE_DIR, "materiales.xlsx")

st.set_page_config(
    page_title="Planes de Lealtad – Grupo Farmatodo",
    page_icon="📋",
    layout="wide",
)

# ── CSS ────────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
[data-testid="stAppViewContainer"] { background: #f4f8fb; }
[data-testid="stSidebar"] { display: none; }
.main-header {
    background: linear-gradient(135deg, #005f73, #0a9396);
    color: white; padding: 1.2rem 2rem; border-radius: 12px;
    margin-bottom: 1.5rem; box-shadow: 0 4px 15px rgba(0,95,115,.3);
}
.main-header h1 { margin: 0; font-size: 1.6rem; }
.main-header p  { margin: .3rem 0 0; opacity: .8; font-size: .95rem; }
.section-card {
    background: white; border: 1px solid #d0e3ea;
    border-radius: 12px; padding: 1.4rem 1.6rem;
    margin-bottom: 1.2rem; box-shadow: 0 1px 4px rgba(0,95,115,.06);
}
.section-title {
    font-size: .75rem; font-weight: 700; letter-spacing: .08em;
    text-transform: uppercase; color: #005f73;
    border-bottom: 2px solid #e9f5f7; padding-bottom: .5rem; margin-bottom: 1rem;
}
.email-box {
    background: #f8fcfd; border: 1.5px solid #0a9396;
    border-radius: 10px; padding: 1.2rem 1.5rem; margin-top: 1rem;
    font-family: monospace; white-space: pre-wrap; font-size: .88rem;
    color: #1a2e35;
}
.found-badge   { color: #06d6a0; font-weight: 700; }
.missing-badge { color: #ef476f; font-weight: 700; }
.stButton>button {
    border-radius: 8px !important; font-weight: 600 !important;
    transition: all .2s !important;
}
</style>
""", unsafe_allow_html=True)

# ── Header ─────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="main-header">
  <h1>📋 Planes de Lealtad</h1>
  <p>Grupo Farmatodo · Generador de convenios y correos</p>
</div>
""", unsafe_allow_html=True)


# ── Load catalog (cached) ──────────────────────────────────────────────────────
@st.cache_data(show_spinner="Cargando catálogo de materiales…")
def load_catalog():
    wb = openpyxl.load_workbook(MATERIALES, read_only=True, data_only=True)
    ws = wb["Hoja1"]
    rows = list(ws.iter_rows(values_only=True))
    sku_map = {}
    lab_map = {}
    for row in rows[1:]:
        sap, sku, desc, id_lab, lab = row
        if sku is None:
            continue
        sku_s = str(int(sku)) if isinstance(sku, float) else str(sku)
        sap_s = str(int(sap)) if sap and isinstance(sap, (int, float)) else (str(sap) if sap else "-")
        entry = {"sap": sap_s, "descripcion": str(desc) if desc else "-",
                 "id_lab": str(int(id_lab)) if id_lab else "-", "lab": str(lab) if lab else "-"}
        sku_map[sku_s] = entry
        if id_lab:
            id_s = str(int(id_lab)) if isinstance(id_lab, float) else str(id_lab)
            if id_s not in lab_map:
                lab_map[id_s] = str(lab) if lab else "-"
    wb.close()
    return sku_map, lab_map

SKU_MAP, LAB_MAP = load_catalog()


# ── Helpers ────────────────────────────────────────────────────────────────────
def format_fecha_es(d: date) -> str:
    months = ["enero","febrero","marzo","abril","mayo","junio",
              "julio","agosto","septiembre","octubre","noviembre","diciembre"]
    return f"{d.day} de {months[d.month-1]}, {d.year}"

def format_vigencia(d: date) -> str:
    months = ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"]
    return f"{d.day:02d}/{months[d.month-1]}/{str(d.year)[2:]}"

def replace_yellow_runs(paragraph, new_text):
    yellow = [r for r in paragraph.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    if not yellow:
        return
    yellow[0].text = new_text
    yellow[0].font.highlight_color = None
    for r in yellow[1:]:
        r.text = ""
        r.font.highlight_color = None

def build_document(fields: dict, productos: list) -> Document:
    doc = Document(PLANTILLA)

    inicio_fmt = format_vigencia(fields["inicio_vigencia"])
    fin_fmt    = format_vigencia(fields["fin_vigencia"])

    # Para 2 – fecha
    replace_yellow_runs(doc.paragraphs[2], fields["fecha_doc"])

    # Para 4 – id_lab, nombre_lab, inicio, fin
    p4 = doc.paragraphs[4]
    yellow4 = [r for r in p4.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    inicio_parts = inicio_fmt.split("/")
    fin_parts    = fin_fmt.split("/")
    all_vals = [fields["id_lab"], fields["nombre_lab"]] + inicio_parts + fin_parts
    for idx, run in enumerate(yellow4):
        run.text = all_vals[idx] if idx < len(all_vals) else ""
        run.font.highlight_color = None

    # Table 0 – forma recuperación
    for p in doc.tables[0].rows[0].cells[0].paragraphs:
        replace_yellow_runs(p, fields["forma_recuperacion"])

    # Table 1 – KAM, puesto, lab
    t1 = doc.tables[1]
    for p in t1.rows[1].cells[1].paragraphs:
        replace_yellow_runs(p, fields["nombre_kam"])
    for p in t1.rows[2].cells[1].paragraphs:
        replace_yellow_runs(p, fields["puesto"])
    for p in t1.rows[3].cells[1].paragraphs:
        replace_yellow_runs(p, fields["nombre_lab"])

    # Table 2 – productos
    t2 = doc.tables[2]
    for row in t2.rows[1:]:
        row._tr.getparent().remove(row._tr)

    vigencia_cell = f"{inicio_fmt} – {fin_fmt}"
    for prod in productos:
        row = t2.add_row()
        for i, val in enumerate([
            str(prod.get("SAP", "-")),
            str(prod.get("SKU", "-")),
            str(prod.get("Producto", "-")),
            str(prod.get("Mecánica", "-")),
            str(prod.get("Límites", "-")),
            vigencia_cell,
        ]):
            row.cells[i].text = val

    return doc

def doc_to_pdf_bytes(doc: Document) -> bytes | None:
    """Returns PDF bytes or None if conversion fails."""
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "convenio.docx")
        pdf_path  = os.path.join(tmpdir, "convenio.pdf")
        doc.save(docx_path)

        # Try LibreOffice
        for lo in ["libreoffice", "soffice"]:
            if subprocess.run(["which", lo], capture_output=True).returncode == 0:
                lo_profile = os.path.join(tmpdir, "lo_profile")
                subprocess.run(
                    [lo, "--headless",
                     f"-env:UserInstallation=file://{lo_profile}",
                     "--convert-to", "pdf", "--outdir", tmpdir, docx_path],
                    capture_output=True, timeout=60
                )
                expected = docx_path.replace(".docx", ".pdf")
                if os.path.exists(expected) and os.path.getsize(expected) > 0:
                    with open(expected, "rb") as f:
                        return f.read()

        # Fallback: mammoth + weasyprint
        try:
            import mammoth
            from weasyprint import HTML
            with open(docx_path, "rb") as f:
                result = mammoth.convert_to_html(f)
            html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"/>
<style>
body{{font-family:Arial,sans-serif;font-size:11pt;margin:2cm;line-height:1.5}}
table{{border-collapse:collapse;width:100%;margin:1em 0;font-size:10pt}}
td,th{{border:1px solid #999;padding:5px 8px}}
th{{background:#2c5f6e;color:white;font-weight:bold}}
p{{margin:.5em 0}}
@page{{margin:2cm}}
</style></head><body>{result.value}</body></html>"""
            return HTML(string=html).write_pdf()
        except Exception:
            return None

def docx_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Session state init ─────────────────────────────────────────────────────────
if "productos_df" not in st.session_state:
    st.session_state.productos_df = pd.DataFrame(
        columns=["SAP", "SKU", "Producto", "Mecánica", "Límites"]
    )
if "doc_generado" not in st.session_state:
    st.session_state.doc_generado = None
if "pdf_bytes" not in st.session_state:
    st.session_state.pdf_bytes = None
if "fields" not in st.session_state:
    st.session_state.fields = {}


# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 1 – Datos generales
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">1 · Datos generales</div>', unsafe_allow_html=True)
c1, c2, c3 = st.columns(3)
accion          = c1.selectbox("Acción *", ["RENOVACION", "NUEVO PLAN"])
nombre_doc      = c2.text_input("Nombre del documento *", placeholder="ej. CONVENIO_ULTRA_2026")
fecha_doc_date  = c3.date_input("Fecha del documento *", value=date.today())
st.markdown('</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 2 – Laboratorio
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">2 · Laboratorio</div>', unsafe_allow_html=True)
c1, c2 = st.columns(2)
id_lab = c1.text_input("ID Laboratorio *", placeholder="ej. 1000660")
nombre_lab = LAB_MAP.get(id_lab.strip(), "") if id_lab.strip() else ""
if id_lab.strip():
    if nombre_lab:
        c2.text_input("Nombre del laboratorio", value=nombre_lab, disabled=True)
    else:
        c2.text_input("Nombre del laboratorio", value="", disabled=True)
        st.warning("⚠ ID de laboratorio no encontrado en el catálogo.")
else:
    c2.text_input("Nombre del laboratorio", value="", disabled=True)
st.markdown('</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 3 – Vigencia y condiciones
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">3 · Vigencia y condiciones</div>', unsafe_allow_html=True)
c1, c2, c3 = st.columns(3)
inicio_vigencia = c1.date_input("Inicio de vigencia *", value=date.today())
fin_vigencia    = c2.date_input("Fin de vigencia *", value=date(date.today().year, 12, 31))
forma_rec       = c3.selectbox("Forma de recuperación *", ["OC COSTO 0", "FACTURA", "NC"])
st.markdown('</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 4 – Firmante
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">4 · Firmante del laboratorio</div>', unsafe_allow_html=True)
c1, c2 = st.columns(2)
nombre_kam = c1.text_input("Nombre del KAM *", placeholder="ej. Laura Edith Castellanos")
puesto     = c2.text_input("Puesto *", placeholder="ej. KAM")
st.markdown('</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 5 – Productos
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">5 · Productos</div>', unsafe_allow_html=True)

c1, c2 = st.columns([2, 1])
sku_input = c1.text_area(
    "Pegar SKUs (uno por línea, o separados por comas/espacios)",
    height=110,
    placeholder="7502216792289\n7502216796737\n7502216796836",
)

with c2:
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("🔍 Buscar productos", use_container_width=True, type="primary"):
        if sku_input.strip():
            skus = [s.strip() for s in sku_input.replace(",", "\n").replace(";", "\n").split() if s.strip()]
            rows = []
            found_count = 0
            for sku in skus:
                entry = SKU_MAP.get(sku)
                if entry:
                    rows.append({"SAP": entry["sap"], "SKU": sku, "Producto": entry["descripcion"],
                                 "Mecánica": "3+1", "Límites": "Sin Límites"})
                    found_count += 1
                else:
                    rows.append({"SAP": "-", "SKU": sku, "Producto": "-",
                                 "Mecánica": "3+1", "Límites": "Sin Límites"})
            st.session_state.productos_df = pd.DataFrame(rows)
            st.success(f"✅ {found_count} de {len(skus)} SKUs encontrados.")
        else:
            st.warning("Pega al menos un SKU primero.")

    if st.button("＋ Agregar fila vacía", use_container_width=True):
        empty = pd.DataFrame([{"SAP": "-", "SKU": "", "Producto": "-", "Mecánica": "3+1", "Límites": "Sin Límites"}])
        st.session_state.productos_df = pd.concat(
            [st.session_state.productos_df, empty], ignore_index=True
        )

    if st.button("🗑 Limpiar tabla", use_container_width=True):
        st.session_state.productos_df = pd.DataFrame(
            columns=["SAP", "SKU", "Producto", "Mecánica", "Límites"]
        )

# Tabla editable
if not st.session_state.productos_df.empty:
    st.markdown("**Revisa y edita los datos antes de generar:**")
    edited_df = st.data_editor(
        st.session_state.productos_df,
        use_container_width=True,
        num_rows="dynamic",
        column_config={
            "SAP":      st.column_config.TextColumn("SAP", width="small"),
            "SKU":      st.column_config.TextColumn("SKU", width="medium"),
            "Producto": st.column_config.TextColumn("Producto", width="large"),
            "Mecánica": st.column_config.TextColumn("Mecánica", width="small"),
            "Límites":  st.column_config.TextColumn("Límites", width="medium"),
        },
        hide_index=True,
        key="tabla_productos",
    )
    st.session_state.productos_df = edited_df
else:
    st.info("Pega SKUs arriba y haz clic en **Buscar productos**, o agrega filas manualmente.")

st.markdown('</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# SECCIÓN 6 – Correo
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="section-card"><div class="section-title">6 · Correo</div>', unsafe_allow_html=True)
c1, c2 = st.columns(2)
correo_to = c1.text_input("Para (destinatario) *", placeholder="kam@laboratorio.com")
correo_cc = c2.text_input("CC (con copia)", placeholder="copia1@lab.com, copia2@lab.com")
st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# GENERAR CONVENIO
# ══════════════════════════════════════════════════════════════════════════════
st.divider()

# Validaciones
def validar():
    errores = []
    if not nombre_lab:
        errores.append("ID de laboratorio no encontrado en catálogo.")
    if not nombre_kam.strip():
        errores.append("Ingresa el nombre del KAM.")
    if st.session_state.productos_df.empty:
        errores.append("Agrega al menos un producto.")
    return errores

col_gen, col_email = st.columns(2)

with col_gen:
    if st.button("📄 Generar convenio y vista previa", type="primary", use_container_width=True):
        errores = validar()
        if errores:
            for e in errores:
                st.error(f"⚠ {e}")
        else:
            fields = {
                "accion": accion,
                "nombre_documento": nombre_doc,
                "fecha_doc": format_fecha_es(fecha_doc_date),
                "id_lab": id_lab.strip(),
                "nombre_lab": nombre_lab,
                "inicio_vigencia": inicio_vigencia,
                "fin_vigencia": fin_vigencia,
                "forma_recuperacion": forma_rec,
                "nombre_kam": nombre_kam.strip(),
                "puesto": puesto.strip(),
            }
            st.session_state.fields = fields
            with st.spinner("Generando documento…"):
                productos = st.session_state.productos_df.to_dict("records")
                doc = build_document(fields, productos)
                st.session_state.doc_generado = docx_bytes(doc)
                doc2 = build_document(fields, productos)  # fresh for PDF
                st.session_state.pdf_bytes = doc_to_pdf_bytes(doc2)
            st.success("✅ Documento generado.")

# ── Vista previa ───────────────────────────────────────────────────────────────
if st.session_state.pdf_bytes or st.session_state.doc_generado:
    st.markdown('<div class="section-card"><div class="section-title">Vista previa del convenio</div>', unsafe_allow_html=True)

    dl1, dl2 = st.columns(2)
    nombre_archivo = st.session_state.fields.get("nombre_documento", "convenio").replace(" ", "_")

    if st.session_state.doc_generado:
        dl1.download_button(
            "⬇ Descargar Word (.docx)",
            data=st.session_state.doc_generado,
            file_name=f"{nombre_archivo}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    if st.session_state.pdf_bytes:
        dl2.download_button(
            "⬇ Descargar PDF",
            data=st.session_state.pdf_bytes,
            file_name=f"{nombre_archivo}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
        # Inline PDF viewer
        b64 = base64.b64encode(st.session_state.pdf_bytes).decode()
        st.markdown(
            f'<iframe src="data:application/pdf;base64,{b64}" width="100%" height="600px" style="border:1px solid #d0e3ea;border-radius:8px;"></iframe>',
            unsafe_allow_html=True,
        )
    else:
        st.info("PDF no disponible en este entorno — descarga el Word para revisar.")

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# PREPARAR CORREO
# ══════════════════════════════════════════════════════════════════════════════
with col_email:
    if st.button("✉ Preparar correo", use_container_width=True):
        errores = validar()
        if errores:
            for e in errores:
                st.error(f"⚠ {e}")
        else:
            st.session_state.mostrar_correo = True

if st.session_state.get("mostrar_correo"):
    primer_nombre = nombre_kam.strip().split()[0] if nombre_kam.strip() else "equipo"
    fecha_str = format_fecha_es(fecha_doc_date)

    if accion == "RENOVACION":
        subject = f"SOLICITUD DE RENOVACIÓN PLAN DE LEALTAD - {nombre_lab} {fecha_str}"
        body = f"""Buen día, {primer_nombre}, ¿cómo estás?

Nos gustaría iniciar el proceso de renovación del Plan de Lealtad que mantenemos con ustedes. Para nosotros es muy importante continuar trabajando de manera conjunta, por lo que agradeceríamos su apoyo para actualizar la carta correspondiente y revisar si será posible mantener las condiciones actuales o si consideran necesario proponer algún ajuste.

Quedamos atentos y en espera de tu apoyo con la revisión y firma de la carta de renovación.

Agradecemos mucho su apoyo y el excelente trabajo que hemos realizado durante este periodo."""
    else:
        subject = f"CARTA PLAN DE LEALTAD {nombre_lab} - GRUPO FT 2026"
        body = f"""Buen día, {primer_nombre}, esperando se encuentren con bien,

Les comparto la carta convenio con los productos revisados para el plan exclusivo de Farmatodo.

¿Me puedes apoyar con la firma del mismo? Por favor

Les agradecemos el apoyo y estamos seguros que tendremos excelentes resultados con su incorporación al Programa.

Cualquier duda o comentario quedamos al pendiente."""

    st.markdown('<div class="section-card"><div class="section-title">✉ Correo generado — edita si necesitas</div>', unsafe_allow_html=True)

    subject_edit = st.text_input("Asunto", value=subject, key="email_subject")
    body_edit    = st.text_area("Cuerpo", value=body, height=220, key="email_body")

    # Mailto link
    import urllib.parse
    params = {"subject": subject_edit, "body": body_edit}
    if correo_cc.strip():
        params["cc"] = correo_cc.strip()
    mailto = f"mailto:{correo_to}?{urllib.parse.urlencode(params)}"
    st.markdown(
        f'<a href="{mailto}" style="display:inline-block;background:#005f73;color:white;padding:.6rem 1.4rem;border-radius:8px;font-weight:600;text-decoration:none;margin-top:.5rem;">📧 Abrir en Outlook / correo</a>',
        unsafe_allow_html=True,
    )
    st.markdown("</div>", unsafe_allow_html=True)
